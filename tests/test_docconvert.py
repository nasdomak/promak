"""The document converter, on tiny documents written on the spot."""

from __future__ import annotations

import os
from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("openpyxl")

from promak.core.filejobs import FileJob, FileStage  # noqa: E402
from promak.core.tables import read_csv, read_xlsx, typed  # noqa: E402
from promak.tools.docconvert import engine as e  # noqa: E402

MARKDOWN = """# Plan

Intro with **bold**, *italic* and `code`.

## Steps
- First
- Second
  - inner
1. one
2. two

> Wise words

| Name | Qty |
| --- | --- |
| Apples | 3 |

```
print("hi")
```
"""


def _convert(source: Path, target: str, out: Path, **options):
    job = FileJob(source=source, destination=out)
    e.ConvertBatch(e.ConvertOptions(target=target, **options)).run([job])
    return job


def test_markdown_is_understood():
    blocks = e.read_markdown(MARKDOWN)
    kinds = [b.kind for b in blocks]
    assert kinds == ["heading", "para", "heading", "bullet", "bullet", "bullet", "number", "number",
                     "quote", "table", "code"]
    assert blocks[5].level == 1
    runs = blocks[1].runs
    assert [(r.text, r.bold, r.italic, r.code) for r in runs if r.bold or r.italic or r.code] == [
        ("bold", True, False, False), ("italic", False, True, False), ("code", False, False, True)]
    assert blocks[9].rows == [["Name", "Qty"], ["Apples", "3"]]


def test_markdown_to_word_and_back(tmp_path):
    source = tmp_path / "plan.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    job = _convert(source, e.TO_DOCX, tmp_path / "word")
    assert job.stage is FileStage.DONE, job.error
    docx_file = tmp_path / "word" / "plan.docx"

    back = _convert(docx_file, e.TO_MD, tmp_path / "md")
    markdown = (tmp_path / "md" / "plan.md").read_text(encoding="utf-8")
    assert markdown.startswith("# Plan")
    assert "**bold**" in markdown and "*italic*" in markdown
    assert "- First" in markdown and "  - inner" in markdown
    assert "| Apples | 3 |" in markdown
    assert "```" in markdown
    assert back.stage is FileStage.DONE

    _convert(docx_file, e.TO_TXT, tmp_path / "txt")
    text = (tmp_path / "txt" / "plan.txt").read_text(encoding="utf-8")
    assert text.startswith("PLAN") and "1. one\n2. two" in text

    _convert(docx_file, e.TO_HTML, tmp_path / "html")
    page = (tmp_path / "html" / "plan.html").read_text(encoding="utf-8")
    assert "<h1>Plan</h1>" in page and "<strong>bold</strong>" in page
    assert "<ul>" in page and "<table>" in page and "<blockquote>" in page


def test_plain_text_to_word(tmp_path):
    source = tmp_path / "letter.txt"
    source.write_text("Dear Anna,\nthank you.\n\nBest regards", encoding="utf-8")
    _convert(source, e.TO_DOCX, tmp_path)
    import docx

    paragraphs = [p.text for p in docx.Document(str(tmp_path / "letter.docx")).paragraphs]
    assert paragraphs == ["Dear Anna, thank you.", "Best regards"]


def test_csv_and_excel_round_trip(tmp_path):
    source = tmp_path / "orders.csv"
    source.write_text("name;amount;zip\nAnna;1234,5;00184\nLuca;7;20100\n", encoding="utf-8")
    rows, dialect = read_csv(source)
    assert dialect.delimiter == ";"
    _convert(source, e.TO_XLSX, tmp_path)
    sheets = read_xlsx(tmp_path / "orders.xlsx")
    assert sheets["orders"] == [["name", "amount", "zip"], ["Anna", 1234.5, "00184"], ["Luca", 7, 20100]]

    _convert(tmp_path / "orders.xlsx", e.TO_CSV, tmp_path / "back", delimiter=";")
    assert (tmp_path / "back" / "orders.csv").read_text(encoding="utf-8-sig").splitlines() == [
        "name;amount;zip", "Anna;1234,5;00184", "Luca;7;20100"]


def test_every_sheet_becomes_a_csv(tmp_path):
    from promak.core.tables import write_xlsx

    book = tmp_path / "book.xlsx"
    write_xlsx(book, {"January": [["a"], [1]], "February": [["b"], [2]]})
    job = _convert(book, e.TO_CSV, tmp_path / "out")
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == ["book - February.csv", "book - January.csv"]
    assert job.stage is FileStage.DONE
    _convert(book, e.TO_CSV, tmp_path / "first", every_sheet=False)
    assert [p.name for p in (tmp_path / "first").iterdir()] == ["book.csv"]


def test_cells_stay_safe():
    assert typed("00184") == "00184"
    assert typed("42") == 42
    assert typed("3,5", comma_decimal=True) == 3.5
    assert typed("3,5") == "3,5"


def test_wrong_pair_is_explained(tmp_path):
    source = tmp_path / "data.csv"
    source.write_text("a,b\n1,2\n", encoding="utf-8")
    job = _convert(source, e.TO_MD, tmp_path)
    assert job.stage is FileStage.FAILED
    assert "cannot become MD" in job.error


@pytest.mark.skipif(e.office_pdf_converter() is None, reason="neither Office nor LibreOffice is installed")
def test_word_to_pdf_with_office(tmp_path):
    source = tmp_path / "plan.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    _convert(source, e.TO_DOCX, tmp_path)
    job = _convert(tmp_path / "plan.docx", e.TO_PDF, tmp_path / "pdf")
    assert job.stage is FileStage.DONE, job.error
    assert (tmp_path / "pdf" / "plan.pdf").read_bytes().startswith(b"%PDF")


def test_pdf_without_office_says_what_to_do(monkeypatch):
    monkeypatch.setattr(e, "office_pdf_converter", lambda: None)
    assert "LibreOffice" in e.ConvertOptions(target=e.TO_PDF).validate()


def test_cli(tmp_path):
    from promak import cli

    source = tmp_path / "plan.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    assert cli.main(["convert", str(source), "--to", "html", "--quiet"]) == 0
    assert (tmp_path / "plan.html").exists()


@pytest.fixture
def qapp(tmp_path, monkeypatch):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    widgets = pytest.importorskip("PySide6.QtWidgets")
    for name in ("APPDATA", "XDG_DATA_HOME"):
        monkeypatch.setenv(name, str(tmp_path / "appdata"))
    import promak.core.config as config_module

    monkeypatch.setattr(config_module, "_config", None)
    return widgets.QApplication.instance() or widgets.QApplication([])


def test_docconvert_screen(qapp, tmp_path):
    from promak.tools.docconvert.panel import DocConvertPanel
    from promak.ui.plan_panel import select

    source = tmp_path / "data.csv"
    source.write_text("a,b\n1,2\n", encoding="utf-8")
    panel = DocConvertPanel()
    try:
        panel.destination_input.setText(str(tmp_path / "out"))
        panel.add_files([source])
        select(panel.target_combo, e.TO_MD)
        assert "No file in the queue" in panel.validate_before_start()
        select(panel.target_combo, e.TO_XLSX)
        assert panel.validate_before_start() is None
        select(panel.target_combo, e.TO_CSV)
        assert not panel.delimiter_combo.isHidden()
    finally:
        panel.deleteLater()
        qapp.processEvents()
