"""Every format other than PDF, read into a :class:`Document`.

=====================  ======================================================
Format                 What is kept
=====================  ======================================================
EPUB (e-books)         chapters in reading order, headings, lists, tables,
                       quotes, pictures, the book's own table of contents
Web pages (HTML)       the same; pictures saved next to the page are copied
Word (DOCX)            headings, lists, tables, bold/italic, pictures
Writer (ODT)           headings, paragraphs, lists, tables, pictures
PowerPoint (PPTX)      one section per slide: title, text, tables, pictures,
                       speaker notes
Excel, CSV             one table per sheet
Pictures               the picture, and the words in it read by the OCR
Text, Markdown         the text as it is
=====================  ======================================================

Only the standard library and what Promak already uses (python-docx,
openpyxl, Pillow) are needed: EPUB, ODT and PPTX files are ZIP folders of
XML and are read directly.  Nothing here imports Qt.
"""

from __future__ import annotations

import html
import logging
import posixpath
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple
from urllib.parse import unquote, urlparse
from xml.etree import ElementTree

from promak.tools.booknotes.model import (
    CODE,
    PAGE,
    PICTURE,
    TABLE,
    TEXT,
    Document,
    NotesError,
    Piece,
    PictureSink,
    clean_inline,
    escape_markdown,
    table_markdown,
)

log = logging.getLogger(__name__)

EPUB_EXTENSIONS = (".epub",)
HTML_EXTENSIONS = (".html", ".htm", ".xhtml")
WORD_EXTENSIONS = (".docx",)
ODT_EXTENSIONS = (".odt",)
SLIDE_EXTENSIONS = (".pptx",)
SHEET_EXTENSIONS = (".xlsx", ".xlsm", ".csv", ".tsv")
PICTURE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif")
TEXT_EXTENSIONS = (".txt", ".md", ".markdown")

#: a sheet longer than this is cut, so one table cannot drown the notes
MAX_SHEET_ROWS = 3000

ImageLoader = Callable[[str], Optional[Tuple[bytes, str]]]


# ================================================================ HTML
class HtmlToPieces(HTMLParser):
    """Turns (X)HTML into pieces: headings, paragraphs, lists, tables, pictures.

    ``load_image(src)`` returns the bytes and the file ending of a picture
    the page points to, or ``None`` when it cannot be found; the picture is
    then saved through ``sink``.
    """

    BLOCKS = {"p", "div", "section", "article", "header", "footer", "aside", "main", "body",
              "figure", "figcaption", "dd", "dt", "address", "center", "nav"}
    SKIP = {"script", "style", "head", "title", "noscript", "svg", "math", "template", "button", "form", "select"}

    VOID = {"img", "image", "br", "hr", "meta", "link", "input", "col", "area", "base", "source", "wbr"}

    def __init__(self, document: Document, sink: Optional[PictureSink], load_image: ImageLoader,
                 where: str, heading_shift: int = 0, styles: Optional["StyleBook"] = None,
                 has_heading_tags: bool = True) -> None:
        super().__init__(convert_charrefs=True)
        self.document = document
        self.styles = styles or StyleBook()
        self.bold_headings = not has_heading_tags
        self.context: List[Tuple[str, float, bool]] = []   # open tags: (tag, letter size, bold)
        self.chunks: List[Tuple[int, float, bool]] = []    # pieces of text of the block: (length, size, bold)
        self.sink = sink
        self.load_image = load_image
        self.where = where
        self.heading_shift = heading_shift
        self.parts: List[str] = []          # the inline text of the current block
        self.skip = 0
        self.heading_level = 0
        self.lists: List[List] = []         # [kind, counter]
        self.quote = 0
        self.pre = 0
        self.table: Optional[List[List[str]]] = None
        self.cell: Optional[List[str]] = None
        self.link: Optional[str] = None
        self.caption_target: Optional[Piece] = None
        self.headings_seen = 0

    # ----------------------------------------------------------- emitting
    def _text_now(self) -> str:
        text = "".join(self.parts)
        self.parts = []
        self._last_chunks, self.chunks = self.chunks, []
        if self.pre:
            return text.strip("\n")
        return re.sub(r"[ \t\r\n ]+", " ", text).replace(" \n", "\n").strip()

    def _flush(self) -> None:
        if self.pre:
            return
        text = self._text_now()
        if not text:
            return
        if self.heading_level:
            self.document.heading(re.sub(r"[*_`]", "", text), self.heading_level + self.heading_shift)
            self.headings_seen += 1
            return
        styled = self._styled_heading(text)
        if styled:
            self.document.heading(re.sub(r"[*_`]", "", text), styled + self.heading_shift)
            self.headings_seen += 1
            return
        if self.lists:
            kind, number = self.lists[-1]
            indent = "  " * (len(self.lists) - 1)
            marker = f"{number}." if kind == "ol" else "-"
            text = f"{indent}{marker} {text}"
        if self.quote:
            text = "\n".join("> " + line for line in text.splitlines())
        piece = Piece(TEXT, text=text)
        self.document.add(piece)

    def _styled_heading(self, text: str) -> int:
        """A short block written all in bigger (or bold) letters is a heading made with styles."""
        chunks = [c for c in getattr(self, "_last_chunks", []) if c[0] > 0]
        if (not chunks or self.lists or self.quote or self.table is not None or len(text) > 150
                or re.search(r"[.,;]$", text.strip("*_ "))):
            return 0
        smallest = min(size for _n, size, _b in chunks)
        if smallest >= 1.15:
            return self.styles.level(smallest)
        all_bold = all(bold for _n, _s, bold in chunks)
        if all_bold and self.bold_headings and len(text) <= 100 and smallest >= 0.98:
            if smallest >= 1.04:
                return self.styles.level(max(1.15, smallest))
            return self.styles.bold_level()
        return 0

    def _add(self, text: str) -> None:
        if self.skip:
            return
        if self.cell is not None:
            self.cell.append(text)
        else:
            self.parts.append(text)

    def _record(self, text: str) -> None:
        """Remember how big and how bold this piece of the block is printed."""
        size = 1.0
        bold = False
        for _tag, tag_size, tag_bold in self.context:
            size = tag_size if tag_size else size
            bold = bold or tag_bold
        self.chunks.append((len(text.strip()), size, bold))

    # -------------------------------------------------------------- tags
    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag in self.SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag not in self.VOID:
            size, bold = self.styles.of(attributes.get("class", ""), attributes.get("style", ""))
            self.context.append((tag, size, bold or tag in ("b", "strong")))
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush()
            self.heading_level = int(tag[1])
        elif tag in self.BLOCKS:
            self._flush()
        elif tag in ("ul", "ol"):
            self._flush()
            self.lists.append([tag, 0])
        elif tag == "li":
            self._flush()
            if self.lists:
                self.lists[-1][1] += 1
        elif tag == "blockquote":
            self._flush()
            self.quote += 1
        elif tag == "pre":
            self._flush()
            self.pre += 1
        elif tag == "br":
            self._add("\n" if self.pre else "  \n")
        elif tag == "hr":
            self._flush()
            self.document.add(Piece(TEXT, text="---"))
        elif tag in ("strong", "b"):
            self._add("**")
        elif tag in ("em", "i"):
            self._add("*")
        elif tag == "code" and not self.pre:
            self._add("`")
        elif tag == "sup":
            self._add("^")
        elif tag == "a":
            href = attributes.get("href") or ""
            self.link = href if href.startswith(("http://", "https://")) else None
            if self.link:
                self._add("[")
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.table.append([])
        elif tag in ("td", "th") and self.table is not None:
            self.cell = []
        elif tag in ("img", "image"):
            src = attributes.get("src") or attributes.get("xlink:href") or attributes.get("href") or ""
            self._picture(src, attributes.get("alt") or attributes.get("title") or "")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        self._end_tag(tag)
        for index in range(len(self.context) - 1, -1, -1):
            if self.context[index][0] == tag:
                del self.context[index:]
                break

    def _end_tag(self, tag) -> None:
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush()
            self.heading_level = 0
        elif tag == "figcaption":
            text = self._text_now()
            if text:
                if self.caption_target is not None and not self.caption_target.picture.alt:
                    self.caption_target.picture.alt = re.sub(r"[*_`]", "", text)
                self.document.add(Piece(TEXT, text=f"*{text.strip('*')}*"))
        elif tag in self.BLOCKS:
            self._flush()
            if tag == "figure":
                self.caption_target = None
        elif tag in ("ul", "ol"):
            self._flush()
            if self.lists:
                self.lists.pop()
        elif tag == "li":
            self._flush()
        elif tag == "blockquote":
            self._flush()
            self.quote = max(0, self.quote - 1)
        elif tag == "pre":
            code = "".join(self.parts).strip("\n")
            self.parts = []
            self.pre = max(0, self.pre - 1)
            if code.strip():
                self.document.add(Piece(CODE, text="```\n" + code + "\n```"))
        elif tag in ("strong", "b"):
            self._add("**")
        elif tag in ("em", "i"):
            self._add("*")
        elif tag == "code" and not self.pre:
            self._add("`")
        elif tag == "a":
            if self.link:
                self._add(f"]({self.link})")
            self.link = None
        elif tag in ("td", "th") and self.table is not None and self.cell is not None:
            if self.table:
                self.table[-1].append(re.sub(r"\s+", " ", "".join(self.cell)).strip())
            self.cell = None
        elif tag == "table" and self.table is not None:
            markdown = table_markdown([row for row in self.table if row])
            self.table = None
            if markdown:
                self.document.add(Piece(TABLE, text=markdown))

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in ("img", "image", "br", "hr"):
            self.handle_endtag(tag)

    def handle_data(self, data):
        if self.skip:
            return
        if not self.pre:
            data = data.replace("<", "&lt;") if re.search(r"<[A-Za-z/]", data) else data
        if self.cell is None:
            self._record(data)
        self._add(data)

    def _picture(self, src: str, alt: str) -> None:
        if not src or self.sink is None:
            return
        if src.startswith(("http://", "https://")):
            # a picture on the internet: link to it, nothing is downloaded
            self._flush()
            self.document.add(Piece(TEXT, text=f"![{clean_inline(alt)}]({src})"))
            return
        loaded = self.load_image(src)
        if loaded is None:
            return
        data, ending = loaded
        saved = self.sink.save_bytes(data, self.where, ending, picture_alt(alt))
        if saved is None:
            return
        self._flush()
        piece = Piece(PICTURE, picture=saved)
        self.document.add(piece)
        self.caption_target = piece

    def close(self) -> None:
        super().close()
        self._flush()


def read_html_text(text: str, document: Document, sink: Optional[PictureSink], load_image: ImageLoader,
                   where: str, heading_shift: int = 0, styles: Optional["StyleBook"] = None) -> int:
    """Read one HTML page into ``document``; returns the number of headings found."""
    book = styles or StyleBook()
    for css in re.findall(r"<style[^>]*>(.*?)</style>", text, re.IGNORECASE | re.DOTALL):
        book.add_css(css)
    tags = bool(re.search(r"<h[1-6][\s>]", text, re.IGNORECASE))
    parser = HtmlToPieces(document, sink, load_image, where, heading_shift, book, tags)
    parser.feed(text)
    parser.close()
    return parser.headings_seen


_PICTURE_FILE = re.compile(r"(^|[/\\])[\w\-. ]+\.(png|jpe?g|gif|webp|svg|bmp|tiff?|emf|wmf)$", re.IGNORECASE)


def picture_alt(text: str) -> str:
    """The description of a picture, unless it is only a file name ("image0001.png")."""
    text = clean_inline(text)
    if _PICTURE_FILE.search(text) or re.fullmatch(r"(image|img|picture|immagine|grafik|figure)?\s*\d*", text, re.I):
        return ""
    return text


class StyleBook:
    """Which style classes print letters bigger or bold, read from the CSS.

    E-books (and LibreOffice) often mark headings with a style, not with a
    heading tag: "a paragraph in 18 pt bold".  The sizes found are ranked,
    so the biggest becomes level 1, the next level 2, and so on.
    """

    KEYWORDS = {"xx-small": 0.6, "x-small": 0.75, "small": 0.89, "medium": 1.0, "large": 1.2,
                "x-large": 1.5, "xx-large": 2.0, "xxx-large": 3.0, "smaller": 0.83, "larger": 1.2}

    def __init__(self) -> None:
        self.classes: Dict[str, Tuple[float, bool]] = {}
        self.base_pt = 12.0
        self._levels: Optional[List[float]] = None

    def add_css(self, css: str) -> None:
        css = re.sub(r"/\*.*?\*/", "", css, flags=re.DOTALL)
        rules = re.findall(r"([^{}]+)\{([^}]*)\}", css)
        for selectors, body in rules:
            if re.search(r"(^|[\s,])(body|p|html)\s*(,|$)", selectors.strip()):
                size = self._points(body)
                if size:
                    self.base_pt = size
        for selectors, body in rules:
            size = self._size(body)
            bold = bool(re.search(r"font-weight\s*:\s*(bold|bolder|[6-9]00)", body, re.IGNORECASE))
            if size is None and not bold:
                continue
            for selector in selectors.split(","):
                names = re.findall(r"\.([A-Za-z_][\w-]*)", selector.strip().split()[-1] if selector.strip() else "")
                for name in names[-1:]:
                    old_size, old_bold = self.classes.get(name, (0.0, False))
                    self.classes[name] = (size or old_size, bold or old_bold)
        self._levels = None

    def _points(self, body: str) -> Optional[float]:
        match = re.search(r"font-size\s*:\s*([\d.]+)\s*(pt|px)", body, re.IGNORECASE)
        if not match:
            return None
        value = float(match.group(1))
        return value if match.group(2).lower() == "pt" else value * 0.75

    def _size(self, body: str) -> Optional[float]:
        match = re.search(r"font-size\s*:\s*([\d.]+)\s*(pt|px|em|rem|%)|font-size\s*:\s*([a-z-]+)", body,
                          re.IGNORECASE)
        if not match:
            return None
        if match.group(3):
            return self.KEYWORDS.get(match.group(3).lower())
        value, unit = float(match.group(1)), match.group(2).lower()
        if unit == "pt":
            return value / self.base_pt
        if unit == "px":
            return value * 0.75 / self.base_pt
        if unit == "%":
            return value / 100.0
        return value

    def of(self, classes: str, style: str = "") -> Tuple[float, bool]:
        size, bold = 0.0, False
        for name in (classes or "").split():
            class_size, class_bold = self.classes.get(name, (0.0, False))
            size = max(size, class_size)
            bold = bold or class_bold
        if style:
            inline = self._size(style)
            size = inline or size
            bold = bold or bool(re.search(r"font-weight\s*:\s*(bold|[6-9]00)", style, re.IGNORECASE))
        return size, bold

    def level(self, size: float) -> int:
        if self._levels is None:
            sizes = sorted({round(s, 2) for s, _b in self.classes.values() if s >= 1.15}, reverse=True)
            self._levels = sizes
        for rank, known in enumerate(self._levels[:4], start=1):
            if size >= known - 0.01:
                return rank
        return min(5, len(self._levels[:4]) + 1) if self._levels else 1

    def bold_level(self) -> int:
        self.level(1.0)
        return min(5, len((self._levels or [])[:4]) + 1)


def _decode(data: bytes) -> str:
    head = data[:2000].decode("ascii", "ignore").lower()
    match = re.search(r'charset=["\']?([a-z0-9_\-]+)', head) or re.search(r'encoding=["\']([a-z0-9_\-]+)', head)
    for encoding in ([match.group(1)] if match else []) + ["utf-8", "cp1252", "latin-1"]:
        try:
            return data.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return data.decode("utf-8", "replace")


def _html_title(text: str) -> str:
    match = re.search(r"<title[^>]*>(.*?)</title>", text, re.IGNORECASE | re.DOTALL)
    return clean_inline(html.unescape(re.sub(r"<[^>]+>", "", match.group(1)))) if match else ""


def read_html(path: Path, sink: Optional[PictureSink]) -> Document:
    text = _decode(Path(path).read_bytes())
    document = Document(kind="html", title=_html_title(text))
    folder = Path(path).parent

    def load(src: str):
        local = folder / unquote(urlparse(src).path)
        try:
            if local.is_file():
                return local.read_bytes(), local.suffix
        except OSError:
            pass
        return None

    read_html_text(text, document, sink, load, "web")
    return document


# ================================================================ EPUB
_OPF = "{http://www.idpf.org/2007/opf}"
_DC = "{http://purl.org/dc/elements/1.1/}"
_CONTAINER = "{urn:oasis:names:tc:opendocument:xmlns:container}"


def _xml(data: bytes):
    try:
        return ElementTree.fromstring(data)
    except ElementTree.ParseError as exc:
        raise NotesError(f"A part of this file is damaged ({exc}).") from exc


def read_epub(path: Path, sink: Optional[PictureSink], report=None) -> Document:
    report = report or (lambda p, d="": None)
    try:
        book = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise NotesError(f"This file cannot be opened as an e-book ({exc}).") from exc
    with book:
        names = set(book.namelist())
        try:
            container = _xml(book.read("META-INF/container.xml"))
            rootfile = container.find(f".//{_CONTAINER}rootfile")
            opf_path = rootfile.get("full-path")
        except (KeyError, AttributeError) as exc:
            raise NotesError("This e-book has no table of contents file (container.xml).") from exc
        if _locked(book, names):
            raise NotesError("This e-book is locked (DRM): it can only be opened in the shop's own app.")
        opf = _xml(book.read(opf_path))
        base = posixpath.dirname(opf_path)
        document = Document(kind="epub")
        metadata = opf.find(f"{_OPF}metadata")
        if metadata is not None:
            title = metadata.find(f"{_DC}title")
            document.title = clean_inline(title.text or "") if title is not None else ""
            creators = [clean_inline(c.text or "") for c in metadata.findall(f"{_DC}creator")]
            document.author = ", ".join(c for c in creators if c)
        manifest: Dict[str, Tuple[str, str, str]] = {}
        for item in opf.iter(f"{_OPF}item"):
            href = posixpath.normpath(posixpath.join(base, unquote(item.get("href", ""))))
            manifest[item.get("id", "")] = (href, item.get("media-type", ""), item.get("properties", ""))
        spine = [manifest[ref.get("idref")][0] for ref in opf.iter(f"{_OPF}itemref")
                 if ref.get("idref") in manifest and "html" in manifest[ref.get("idref")][1]]
        titles = _epub_toc(book, opf, manifest, base)
        styles = StyleBook()
        for href, kind, _props in manifest.values():
            if kind == "text/css" and href in names:
                styles.add_css(_decode(book.read(href)))
        document.pages = len(spine)
        for number, href in enumerate(spine, start=1):
            report(5 + 90 * number / max(1, len(spine)), f"chapter file {number} of {len(spine)}")
            try:
                text = _decode(book.read(href))
            except KeyError:
                continue
            folder = posixpath.dirname(href)

            def load(src: str, folder=folder):
                target = posixpath.normpath(posixpath.join(folder, unquote(urlparse(src).path)))
                if target in names:
                    return book.read(target), posixpath.splitext(target)[1]
                return None

            document.add(Piece(PAGE, page=f"part {number}", page_index=number))
            before = len(document.pieces)
            found = read_html_text(text, document, sink, load, f"c{number:03d}", styles=styles)
            title = titles.get(href)
            if title and not found:
                # chapters marked only with styles: the book's own contents gives the title
                document.pieces.insert(before, Piece("heading", text=title, level=1))
        return document


def _locked(book: zipfile.ZipFile, names) -> bool:
    """True when the chapters themselves are encrypted (fonts alone may be, harmlessly)."""
    if "META-INF/encryption.xml" not in names:
        return False
    try:
        root = _xml(book.read("META-INF/encryption.xml"))
    except NotesError:
        return False
    for reference in root.iter("{http://www.w3.org/2001/04/xmlenc#}CipherReference"):
        if re.search(r"\.(x?html?|xml)$", reference.get("URI", ""), re.IGNORECASE):
            return True
    return False


def _epub_toc(book: zipfile.ZipFile, opf, manifest, base: str) -> Dict[str, str]:
    """The first title the e-book's contents gives to each chapter file."""
    titles: Dict[str, str] = {}
    nav = next((href for href, kind, props in manifest.values() if "nav" in props.split()), None)
    ncx = next((href for href, kind, props in manifest.values() if kind == "application/x-dtbncx+xml"), None)
    try:
        if nav:
            text = _decode(book.read(nav))
            folder = posixpath.dirname(nav)
            for href, label in re.findall(r'<a[^>]+href="([^"#]+)[^"]*"[^>]*>(.*?)</a>', text, re.DOTALL):
                key = posixpath.normpath(posixpath.join(folder, unquote(href)))
                titles.setdefault(key, clean_inline(html.unescape(re.sub(r"<[^>]+>", "", label))))
        elif ncx:
            root = _xml(book.read(ncx))
            folder = posixpath.dirname(ncx)
            ns = "{http://www.daisy.org/z3986/2005/ncx/}"
            for point in root.iter(f"{ns}navPoint"):
                label = point.find(f"{ns}navLabel/{ns}text")
                content = point.find(f"{ns}content")
                if label is None or content is None:
                    continue
                key = posixpath.normpath(posixpath.join(folder, unquote(content.get("src", "").split("#")[0])))
                titles.setdefault(key, clean_inline(label.text or ""))
    except (KeyError, NotesError):
        pass
    return titles


# ================================================================ Word
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_V = "{urn:schemas-microsoft-com:vml}"


def read_docx(path: Path, sink: Optional[PictureSink]) -> Document:
    from promak.tools.docconvert.engine import _paragraph_block, _dedupe_merged, require_docx, to_markdown

    docx = require_docx()
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        word = docx.Document(str(path))
    except Exception as exc:
        raise NotesError(f"This file cannot be read as a Word document ({exc}).") from exc
    document = Document(kind="docx")
    try:
        document.title = clean_inline(word.core_properties.title or "")
        document.author = clean_inline(word.core_properties.author or "")
    except Exception:
        pass
    related = word.part.related_parts
    for element in word.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, word)
            block = _paragraph_block(paragraph)
            if block is not None:
                if block.kind == "heading":
                    document.heading(block.text, block.level)
                else:
                    document.add(Piece(TEXT, text=to_markdown([block]).strip()))
            for rid in _picture_ids(element):
                part = related.get(rid)
                if part is None or sink is None:
                    continue
                ending = posixpath.splitext(str(getattr(part, "partname", "")))[1]
                saved = sink.save_bytes(part.blob, "w", ending, picture_alt(_picture_name(element)))
                if saved is not None:
                    document.add(Piece(PICTURE, picture=saved))
        elif tag == "tbl":
            rows = [[cell.text.strip() for cell in row.cells] for row in Table(element, word).rows]
            markdown = table_markdown(_dedupe_merged(rows))
            if markdown:
                document.add(Piece(TABLE, text=markdown))
    return document


def _picture_ids(element) -> List[str]:
    ids = [blip.get(f"{_R}embed") for blip in element.iter(f"{_A}blip")]
    ids += [image.get(f"{_R}id") for image in element.iter(f"{_V}imagedata")]
    return [rid for rid in ids if rid]


def _picture_name(element) -> str:
    for props in element.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"):
        text = props.get("descr") or props.get("title") or ""
        if text:
            return clean_inline(text)
    return ""


# ================================================================ Writer (ODT)
_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_DRAW = "{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}"
_XLINK = "{http://www.w3.org/1999/xlink}"
_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
_META = "{urn:oasis:names:tc:opendocument:xmlns:meta:1.0}"
_SVG = "{urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0}"


def _odf_text(element) -> str:
    """The words of an ODF element, with its spaces, tabs and line breaks."""
    out = [element.text or ""]
    for child in element:
        tag = child.tag
        if tag == f"{_TEXT}s":
            out.append(" " * int(child.get(f"{_TEXT}c", "1")))
        elif tag == f"{_TEXT}tab":
            out.append("\t")
        elif tag == f"{_TEXT}line-break":
            out.append("  \n")
        elif tag in (f"{_TEXT}note",):
            body = child.find(f"{_TEXT}note-body")
            if body is not None:
                out.append(" (note: " + " ".join(_odf_text(p) for p in body) + ")")
        elif tag != f"{_DRAW}frame":
            out.append(_odf_text(child))
        out.append(child.tail or "")
    return "".join(out)


def read_odt(path: Path, sink: Optional[PictureSink]) -> Document:
    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise NotesError(f"This file cannot be opened as a Writer document ({exc}).") from exc
    with archive:
        document = Document(kind="odt")
        try:
            meta = _xml(archive.read("meta.xml"))
            title = meta.find(f".//{_DC}title")
            creator = meta.find(f".//{_DC}creator") or meta.find(f".//{_META}initial-creator")
            document.title = clean_inline(title.text or "") if title is not None else ""
            document.author = clean_inline(creator.text or "") if creator is not None else ""
        except KeyError:
            pass
        try:
            content = _xml(archive.read("content.xml"))
        except KeyError as exc:
            raise NotesError("This Writer file has no content.") from exc
        names = set(archive.namelist())
        body = content.find(f"{_OFFICE}body/{_OFFICE}text")
        if body is None:
            return document

        def pictures(element) -> None:
            for frame in element.iter(f"{_DRAW}frame"):
                image = frame.find(f"{_DRAW}image")
                if image is None or sink is None:
                    continue
                href = image.get(f"{_XLINK}href", "")
                if href not in names:
                    continue
                title = frame.find(f"{_SVG}title")
                desc = frame.find(f"{_SVG}desc")
                alt = picture_alt((desc.text if desc is not None else "") or (title.text if title is not None else "") or "")
                saved = sink.save_bytes(archive.read(href), "w", posixpath.splitext(href)[1], alt)
                if saved is not None:
                    document.add(Piece(PICTURE, picture=saved))

        def walk(element, depth: int = 0) -> None:
            for child in element:
                tag = child.tag
                if tag == f"{_TEXT}h":
                    document.heading(_odf_text(child), int(child.get(f"{_TEXT}outline-level", "1")))
                    pictures(child)
                elif tag == f"{_TEXT}p":
                    text = _odf_text(child).strip()
                    if text:
                        document.add(Piece(TEXT, text=escape_markdown(text)))
                    pictures(child)
                elif tag == f"{_TEXT}list":
                    walk_list(child, depth)
                elif tag == f"{_TABLE}table":
                    rows = []
                    for row in child.iter(f"{_TABLE}table-row"):
                        cells = []
                        for cell in row:
                            if cell.tag in (f"{_TABLE}table-cell", f"{_TABLE}covered-table-cell"):
                                repeat = min(50, int(cell.get(f"{_TABLE}number-columns-repeated", "1")))
                                value = " ".join(_odf_text(p).strip() for p in cell)
                                cells += [value] * repeat
                        while cells and not cells[-1]:
                            cells.pop()
                        rows.append(cells)
                    markdown = table_markdown(rows)
                    if markdown:
                        document.add(Piece(TABLE, text=markdown))
                elif tag in (f"{_TEXT}section", f"{_TEXT}index-body", f"{_TEXT}table-of-content"):
                    if tag == f"{_TEXT}table-of-content":
                        continue      # the printed contents page: the headings already say it
                    walk(child, depth)

        def walk_list(element, depth: int) -> None:
            for item in element.findall(f"{_TEXT}list-item"):
                for part in item:
                    if part.tag == f"{_TEXT}list":
                        walk_list(part, depth + 1)
                    else:
                        text = _odf_text(part).strip()
                        if text:
                            document.add(Piece(TEXT, text=f"{'  ' * depth}- {escape_markdown(text)}"))
                        pictures(part)

        walk(body)
        return document


# ================================================================ PowerPoint
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _rels(archive: zipfile.ZipFile, part: str) -> Dict[str, Tuple[str, str]]:
    folder, name = posixpath.split(part)
    rels_name = posixpath.join(folder, "_rels", name + ".rels")
    try:
        root = _xml(archive.read(rels_name))
    except KeyError:
        return {}
    out = {}
    for rel in root.iter(f"{_REL}Relationship"):
        target = rel.get("Target", "")
        if rel.get("TargetMode") != "External":
            target = posixpath.normpath(posixpath.join(folder, target))
        out[rel.get("Id", "")] = (target, rel.get("Type", "").rsplit("/", 1)[-1])
    return out


def _slide_paragraphs(shape) -> List[Tuple[int, str]]:
    out = []
    for paragraph in shape.iter(f"{_A}p"):
        text = "".join(node.text or "" for node in paragraph.iter() if node.tag in (f"{_A}t",))
        props = paragraph.find(f"{_A}pPr")
        level = int(props.get("lvl", "0")) if props is not None else 0
        if text.strip():
            out.append((level, clean_inline(text)))
    return out


def read_pptx(path: Path, sink: Optional[PictureSink]) -> Document:
    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise NotesError(f"This file cannot be opened as a PowerPoint file ({exc}).") from exc
    with archive:
        document = Document(kind="pptx")
        try:
            core = _xml(archive.read("docProps/core.xml"))
            title = core.find(f"{_DC}title")
            creator = core.find(f"{_DC}creator")
            document.title = clean_inline(title.text or "") if title is not None else ""
            document.author = clean_inline(creator.text or "") if creator is not None else ""
        except KeyError:
            pass
        try:
            presentation = _xml(archive.read("ppt/presentation.xml"))
        except KeyError as exc:
            raise NotesError("This PowerPoint file has no slides.") from exc
        rels = _rels(archive, "ppt/presentation.xml")
        slides = [rels[s.get(f"{_R}id")][0] for s in presentation.iter(f"{_P}sldId") if s.get(f"{_R}id") in rels]
        document.pages = len(slides)
        for number, slide_part in enumerate(slides, start=1):
            try:
                slide = _xml(archive.read(slide_part))
            except KeyError:
                continue
            slide_rels = _rels(archive, slide_part)
            document.add(Piece(PAGE, page=f"slide {number}", page_index=number))
            title, body, tables, pictures = "", [], [], []
            tree = slide.find(f"{_P}cSld/{_P}spTree")
            for shape in (tree.iter() if tree is not None else []):
                if shape.tag == f"{_P}sp":
                    placeholder = shape.find(f"{_P}nvSpPr/{_P}nvPr/{_P}ph")
                    kind = placeholder.get("type", "") if placeholder is not None else ""
                    paragraphs = _slide_paragraphs(shape)
                    if kind in ("title", "ctrTitle") and not title:
                        title = " ".join(text for _level, text in paragraphs)
                    elif kind in ("sldNum", "dt", "ftr"):
                        continue
                    else:
                        body.extend(paragraphs)
                elif shape.tag == f"{_A}tbl":
                    rows = [[" ".join(t.text or "" for t in cell.iter(f"{_A}t")) for cell in row.findall(f"{_A}tc")]
                            for row in shape.findall(f"{_A}tr")]
                    tables.append(table_markdown(rows))
                elif shape.tag == f"{_P}pic":
                    blip = shape.find(f".//{_A}blip")
                    if blip is not None and blip.get(f"{_R}embed") in slide_rels:
                        name = shape.find(f"{_P}nvPicPr/{_P}cNvPr")
                        alt = picture_alt((name.get("descr") or "") if name is not None else "")
                        pictures.append((slide_rels[blip.get(f"{_R}embed")][0], alt))
            document.heading(f"Slide {number}" + (f": {title}" if title else ""), 1)
            if body:
                lines = [f"{'  ' * level}- {escape_markdown(text)}" if len(body) > 1 else escape_markdown(text)
                         for level, text in body]
                document.add(Piece(TEXT, text="\n".join(lines)))
            for table in tables:
                if table:
                    document.add(Piece(TABLE, text=table))
            for target, alt in pictures:
                if sink is None or target not in archive.namelist():
                    continue
                saved = sink.save_bytes(archive.read(target), f"s{number:03d}", posixpath.splitext(target)[1], alt)
                if saved is not None:
                    document.add(Piece(PICTURE, picture=saved))
            notes = next((target for target, kind in slide_rels.values() if kind == "notesSlide"), None)
            if notes:
                try:
                    notes_xml = _xml(archive.read(notes))
                    words = []
                    for shape in notes_xml.iter(f"{_P}sp"):
                        placeholder = shape.find(f"{_P}nvSpPr/{_P}nvPr/{_P}ph")
                        if placeholder is not None and placeholder.get("type") == "body":
                            words += [text for _level, text in _slide_paragraphs(shape)]
                    if words:
                        document.add(Piece(TEXT, text="> **Speaker notes:** " + " ".join(words)))
                except KeyError:
                    pass
        return document


# ================================================================ sheets
def read_sheets(path: Path) -> Document:
    from promak.core.tables import cell_text, read_table

    try:
        sheets = read_table(path)
    except Exception as exc:
        raise NotesError(f"This table cannot be read ({exc}).") from exc
    document = Document(kind=path.suffix.lstrip(".").lower(), pages=len(sheets))
    for name, rows in sheets.items():
        rows = [[cell_text(cell) for cell in row] for row in rows]
        while rows and not any(rows[-1]):
            rows.pop()
        if len(sheets) > 1:
            document.heading(str(name), 1)
        if not rows:
            document.text("*(empty sheet)*")
            continue
        cut = len(rows) > MAX_SHEET_ROWS
        markdown = table_markdown(rows[:MAX_SHEET_ROWS])
        document.add(Piece(TABLE, text=markdown))
        if cut:
            document.text(f"*Only the first {MAX_SHEET_ROWS} of {len(rows)} rows are shown.*")
            document.notes.append(f"sheet {name}: only the first {MAX_SHEET_ROWS} rows kept")
    return document


# ================================================================ pictures
def read_picture(path: Path, sink: Optional[PictureSink], ocr: bool) -> Document:
    from PIL import Image, ImageOps

    try:
        with Image.open(path) as opened:
            opened.load()
            image = ImageOps.exif_transpose(opened)
    except Exception as exc:
        raise NotesError(f"The picture cannot be opened: {exc}") from exc
    document = Document(kind=path.suffix.lstrip(".").lower(), pages=1)
    if sink is not None:
        saved = sink.save_image(image, "img", alt=path.stem)
        if saved is not None:
            document.add(Piece(PICTURE, picture=saved))
    if ocr:
        try:
            from promak.tools.ocr.engine import read_words, words_to_text

            text = words_to_text(read_words(image, 0.5))
        except Exception as exc:
            document.notes.append(f"the words in the picture were not read ({str(exc).split('.')[0]})")
            text = ""
        for paragraph in re.split(r"\n\s*\n", text):
            if paragraph.strip():
                document.text(escape_markdown(" ".join(paragraph.split())))
        if text:
            document.ocr_pages = 1
    return document


# ================================================================ text
def read_plain(path: Path) -> Document:
    from promak.tools.text.engine import read_text

    text = read_text(path)
    document = Document(kind=path.suffix.lstrip(".").lower())
    if path.suffix.lower() in (".md", ".markdown"):
        body = re.sub(r"^---\n.*?\n---\n", "", text, count=1, flags=re.DOTALL)   # its own front matter
        match = re.search(r"^#\s+(.+)$", body, re.MULTILINE)
        document.title = clean_inline(match.group(1)) if match else ""
        for block in re.split(r"\n\s*\n", body):
            heading = re.fullmatch(r"(#{1,6})\s+(.+)", block.strip())
            if heading:
                document.heading(heading.group(2), len(heading.group(1)))
            elif block.strip():
                document.add(Piece(TEXT, text=block.strip("\n")))
        return document
    for block in re.split(r"\n\s*\n", text):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if lines:
            document.text(escape_markdown(" ".join(lines)))
    return document

