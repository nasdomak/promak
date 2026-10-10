"""Converting documents and spreadsheets between formats.

=================  ====================================================
From               To
=================  ====================================================
Word (DOCX)        plain text, Markdown, web page (HTML), PDF*
Markdown, text     Word (DOCX), web page (HTML)
Excel (XLSX)       CSV (one file per sheet, or only the first), PDF*
CSV                Excel (XLSX)
Office files       PDF* (also DOC, ODT, RTF, XLS, ODS, PPT, PPTX, ODP)
=================  ====================================================

\\* PDF needs Microsoft Office (Windows) or LibreOffice (any system) on the
computer: they draw the pages exactly as they look.  Promak only asks them
to do it, without showing a window; when neither is there it says so.

Word files are read and written with python-docx (MIT), workbooks with
openpyxl (MIT).  Headings, paragraphs, bold and italic, lists, quotes,
code and tables are carried over; pictures inside a document are not.
Nothing here imports Qt.
"""

from __future__ import annotations

import html
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from promak.core.batch import BatchEngine
from promak.core.dependencies import SUBPROCESS_QUIET
from promak.core.filejobs import FileJob
from promak.core.imaging import ImageToolError, human_size
from promak.core.paths import safe_filename
from promak.core.tables import read_csv, read_xlsx, typed, write_csv, write_xlsx
from promak.tools.pdf.engine import output_path

log = logging.getLogger(__name__)

WORD_EXTENSIONS = (".docx",)
TEXT_EXTENSIONS = (".md", ".markdown", ".txt")
SHEET_EXTENSIONS = (".xlsx", ".xlsm")
CSV_EXTENSIONS = (".csv", ".tsv")
OFFICE_EXTENSIONS = (".docx", ".doc", ".odt", ".rtf", ".xlsx", ".xls", ".ods", ".pptx", ".ppt", ".odp")
ACCEPTED_EXTENSIONS = tuple(dict.fromkeys(WORD_EXTENSIONS + TEXT_EXTENSIONS + SHEET_EXTENSIONS
                                          + CSV_EXTENSIONS + OFFICE_EXTENSIONS))

TO_TXT, TO_MD, TO_HTML, TO_DOCX, TO_CSV, TO_XLSX, TO_PDF = "txt", "md", "html", "docx", "csv", "xlsx", "pdf"
TARGETS = [
    ("Plain text (TXT)  - from Word", TO_TXT),
    ("Markdown (MD)  - from Word", TO_MD),
    ("Web page (HTML)  - from Word, Markdown or text", TO_HTML),
    ("Word (DOCX)  - from Markdown or text", TO_DOCX),
    ("CSV  - from Excel", TO_CSV),
    ("Excel (XLSX)  - from CSV", TO_XLSX),
    ("PDF  - from Word, Excel, PowerPoint... (needs Office or LibreOffice)", TO_PDF),
]

#: what each target can be made from
SOURCES = {
    TO_TXT: WORD_EXTENSIONS,
    TO_MD: WORD_EXTENSIONS,
    TO_HTML: WORD_EXTENSIONS + TEXT_EXTENSIONS,
    TO_DOCX: TEXT_EXTENSIONS,
    TO_CSV: SHEET_EXTENSIONS,
    TO_XLSX: CSV_EXTENSIONS,
    TO_PDF: OFFICE_EXTENSIONS,
}

DELIMITERS = [("Comma  ,", ","), ("Semicolon  ;  (Excel where the decimal mark is a comma)", ";"), ("Tab", "\t")]


class ConvertError(ImageToolError):
    """A conversion problem told in plain words."""


@dataclass
class ConvertOptions:
    target: str = TO_TXT
    every_sheet: bool = True        # XLSX -> CSV: one CSV per sheet
    delimiter: str = ","            # for the CSV written
    suffix: str = ""
    overwrite: bool = False

    def validate(self) -> Optional[str]:
        if self.target not in SOURCES:
            return "Choose the format to convert to."
        if self.target == TO_PDF and office_pdf_converter() is None:
            return NO_OFFICE_MESSAGE
        return None


# ============================================================ the document
@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    link: str = ""


@dataclass
class Block:
    """One paragraph-sized piece of a document."""

    kind: str                     # heading, para, bullet, number, quote, code, table
    runs: List[Run] = field(default_factory=list)
    level: int = 0                # heading level 1-6, list depth 0..
    rows: List[List[str]] = field(default_factory=list)   # tables

    @property
    def text(self) -> str:
        return "".join(run.text for run in self.runs)


# ------------------------------------------------------------ reading DOCX
def require_docx():
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - depends on the machine
        raise ConvertError("python-docx is missing, so Word files cannot be read or written. "
                           "Run install_windows.bat again, or:  pip install -U python-docx") from exc
    return docx


def read_docx(path: Path) -> List[Block]:
    docx = require_docx()
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ConvertError(f"This file cannot be read as a Word document ({exc}).") from exc
    blocks: List[Block] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            block = _paragraph_block(Paragraph(element, document))
            if block is not None:
                blocks.append(block)
        elif tag == "tbl":
            table = Table(element, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(Block("table", rows=_dedupe_merged(rows)))
    return blocks


def _dedupe_merged(rows: List[List[str]]) -> List[List[str]]:
    """python-docx repeats a merged cell in every column it spans; keep it once."""
    out = []
    for row in rows:
        clean: List[str] = []
        for cell in row:
            clean.append("" if clean and cell and cell == clean[-1] else cell)
        out.append(clean)
    return out


def _paragraph_block(paragraph) -> Optional[Block]:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    runs = _runs(paragraph)
    if not "".join(r.text for r in runs).strip():
        return None
    lowered = style.lower()
    match = re.match(r"heading (\d)", lowered)
    if lowered == "title":
        return Block("heading", runs, level=1)
    if match:
        return Block("heading", runs, level=min(6, int(match.group(1))))
    numbered = paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
    if "list number" in lowered:
        return Block("number", runs, level=_list_level(paragraph, lowered))
    if "list" in lowered or numbered:
        return Block("bullet", runs, level=_list_level(paragraph, lowered))
    if "quote" in lowered:
        return Block("quote", runs)
    if "code" in lowered or "html preformatted" in lowered or (
            lowered == "no spacing" and all(run.code for run in runs)):
        return Block("code", [Run(paragraph.text)])
    return Block("para", runs)


def _list_level(paragraph, style: str) -> int:
    match = re.search(r"(\d)$", style)
    if match and int(match.group(1)) > 1:
        return int(match.group(1)) - 1
    try:
        return int(paragraph._p.pPr.numPr.ilvl.val)
    except AttributeError:
        return 0


def _runs(paragraph) -> List[Run]:
    runs: List[Run] = []
    items = paragraph.iter_inner_content() if hasattr(paragraph, "iter_inner_content") else paragraph.runs
    for item in items:
        if hasattr(item, "runs") and hasattr(item, "address"):   # a hyperlink
            url = getattr(item, "url", "") or getattr(item, "address", "")
            runs.append(Run(item.text, link=url))
            continue
        text = item.text
        if not text:
            continue
        font_name = (item.font.name or "") if item.font is not None else ""
        runs.append(Run(text, bold=bool(item.bold), italic=bool(item.italic),
                        code=font_name.lower() in ("consolas", "courier new", "courier")))
    return _merge_runs(runs)


def _merge_runs(runs: List[Run]) -> List[Run]:
    merged: List[Run] = []
    for run in runs:
        if merged and (merged[-1].bold, merged[-1].italic, merged[-1].code, merged[-1].link) == \
                (run.bold, run.italic, run.code, run.link):
            merged[-1].text += run.text
        else:
            merged.append(Run(run.text, run.bold, run.italic, run.code, run.link))
    return merged


# -------------------------------------------------------- reading Markdown
_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_NUMBER = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE = re.compile(r"(\*\*|__)(.+?)\1|(\*|_)(?!\s)(.+?)(?<!\s)\3|`([^`]+)`|\[([^\]]+)\]\(([^)\s]+)\)")


def parse_inline(text: str) -> List[Run]:
    """``**bold**``, ``*italic*``, ``code`` and ``[links](url)`` into runs."""
    runs: List[Run] = []
    position = 0
    for match in _INLINE.finditer(text):
        if match.start() > position:
            runs.append(Run(text[position:match.start()]))
        if match.group(2) is not None:
            runs.extend(Run(r.text, bold=True, italic=r.italic, code=r.code, link=r.link)
                        for r in parse_inline(match.group(2)))
        elif match.group(4) is not None:
            runs.extend(Run(r.text, bold=r.bold, italic=True, code=r.code, link=r.link)
                        for r in parse_inline(match.group(4)))
        elif match.group(5) is not None:
            runs.append(Run(match.group(5), code=True))
        else:
            runs.append(Run(match.group(6), link=match.group(7)))
        position = match.end()
    if position < len(text):
        runs.append(Run(text[position:]))
    return _merge_runs(runs) if runs else []


def read_markdown(text: str) -> List[Block]:
    blocks: List[Block] = []
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    paragraph: List[str] = []
    index = 0

    def flush() -> None:
        if paragraph:
            blocks.append(Block("para", parse_inline(" ".join(line.strip() for line in paragraph))))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            code: List[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append(Block("code", [Run("\n".join(code))]))
            index += 1
            continue
        if not stripped:
            flush()
        elif _HEADING.match(stripped):
            flush()
            hashes, title = _HEADING.match(stripped).groups()
            blocks.append(Block("heading", parse_inline(title), level=len(hashes)))
        elif _RULE.match(stripped):
            flush()
        elif stripped.startswith("|") and index + 1 < len(lines) and _TABLE_SEPARATOR.match(lines[index + 1]):
            flush()
            rows = [_table_cells(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_table_cells(lines[index].strip()))
                index += 1
            blocks.append(Block("table", rows=rows))
            continue
        elif _BULLET.match(line):
            flush()
            indent, body = _BULLET.match(line).groups()
            blocks.append(Block("bullet", parse_inline(body), level=len(indent.replace("\t", "    ")) // 2))
        elif _NUMBER.match(line):
            flush()
            indent, body = _NUMBER.match(line).groups()
            blocks.append(Block("number", parse_inline(body), level=len(indent.replace("\t", "    ")) // 2))
        elif stripped.startswith(">"):
            flush()
            blocks.append(Block("quote", parse_inline(stripped.lstrip("> ").strip())))
        else:
            paragraph.append(line)
        index += 1
    flush()
    return blocks


def _table_cells(line: str) -> List[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def read_plain_text(text: str) -> List[Block]:
    """Paragraphs are separated by empty lines; nothing else is interpreted."""
    paragraphs = re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip())
    return [Block("para", [Run(" ".join(p.split("\n")).strip())]) for p in paragraphs if p.strip()]


# ------------------------------------------------------------------ writing
def to_text(blocks: Sequence[Block]) -> str:
    out: List[str] = []
    counters: dict = {}
    for block in blocks:
        if block.kind != "number":
            counters.clear()
        if block.kind == "table":
            out.append("\n".join("\t".join(row) for row in block.rows))
        elif block.kind == "bullet":
            out.append("  " * block.level + "- " + block.text)
        elif block.kind == "number":
            counters[block.level] = counters.get(block.level, 0) + 1
            out.append("  " * block.level + f"{counters[block.level]}. " + block.text)
        elif block.kind == "heading":
            out.append(block.text.upper() if block.level == 1 else block.text)
        else:
            out.append(block.text)
    return _join_blocks(blocks, out)


def _join_blocks(blocks: Sequence[Block], pieces: List[str]) -> str:
    """List items sit on consecutive lines; everything else has an empty line between."""
    text = ""
    for index, (block, piece) in enumerate(zip(blocks, pieces)):
        if index:
            joined = block.kind in ("bullet", "number") and blocks[index - 1].kind in ("bullet", "number")
            text += "\n" if joined else "\n\n"
        text += piece
    return text.strip() + "\n"


def _md_runs(runs: Sequence[Run]) -> str:
    out = []
    for run in runs:
        text = run.text
        if not text.strip():
            out.append(text)
            continue
        lead = text[: len(text) - len(text.lstrip())]
        tail = text[len(text.rstrip()):]
        core = text.strip()
        if run.code:
            core = f"`{core}`"
        if run.italic:
            core = f"*{core}*"
        if run.bold:
            core = f"**{core}**"
        if run.link:
            core = f"[{core}]({run.link})"
        out.append(lead + core + tail)
    return "".join(out)


def to_markdown(blocks: Sequence[Block]) -> str:
    out: List[str] = []
    for block in blocks:
        if block.kind == "heading":
            out.append("#" * block.level + " " + _md_runs(block.runs))
        elif block.kind == "bullet":
            out.append("  " * block.level + "- " + _md_runs(block.runs))
        elif block.kind == "number":
            out.append("  " * block.level + "1. " + _md_runs(block.runs))
        elif block.kind == "quote":
            out.append("> " + _md_runs(block.runs))
        elif block.kind == "code":
            out.append("```\n" + block.text + "\n```")
        elif block.kind == "table":
            width = max(len(row) for row in block.rows)
            rows = [row + [""] * (width - len(row)) for row in block.rows]
            lines = ["| " + " | ".join(cell.replace("|", "\\|") for cell in rows[0]) + " |",
                     "|" + "|".join(" --- " for _ in range(width)) + "|"]
            lines += ["| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |" for row in rows[1:]]
            out.append("\n".join(lines))
        else:
            out.append(_md_runs(block.runs))
    return _join_blocks(blocks, out)


def _html_runs(runs: Sequence[Run]) -> str:
    out = []
    for run in runs:
        text = html.escape(run.text)
        if run.code:
            text = f"<code>{text}</code>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.link:
            text = f'<a href="{html.escape(run.link, quote=True)}">{text}</a>'
        out.append(text)
    return "".join(out)


def to_html(blocks: Sequence[Block], title: str) -> str:
    body: List[str] = []
    open_list: List[str] = []

    def close_lists(depth: int = 0) -> None:
        while len(open_list) > depth:
            body.append(f"</{open_list.pop()}>")

    for block in blocks:
        if block.kind in ("bullet", "number"):
            tag = "ul" if block.kind == "bullet" else "ol"
            depth = block.level + 1
            close_lists(depth)
            if len(open_list) == depth and open_list[-1] != tag:
                close_lists(depth - 1)
            while len(open_list) < depth:
                body.append(f"<{tag}>")
                open_list.append(tag)
            body.append(f"<li>{_html_runs(block.runs)}</li>")
            continue
        close_lists()
        if block.kind == "heading":
            body.append(f"<h{block.level}>{_html_runs(block.runs)}</h{block.level}>")
        elif block.kind == "quote":
            body.append(f"<blockquote>{_html_runs(block.runs)}</blockquote>")
        elif block.kind == "code":
            body.append(f"<pre><code>{html.escape(block.text)}</code></pre>")
        elif block.kind == "table":
            rows = block.rows
            head = "".join(f"<th>{html.escape(cell)}</th>" for cell in rows[0])
            rest = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>" for row in rows[1:])
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rest}</tbody></table>")
        else:
            body.append(f"<p>{_html_runs(block.runs)}</p>")
    close_lists()
    return ("<!doctype html>\n<html><head><meta charset=\"utf-8\">"
            f"<title>{html.escape(title)}</title>"
            "<style>body{font-family:system-ui,sans-serif;max-width:46em;margin:3em auto;padding:0 1em;"
            "line-height:1.6;color:#222}table{border-collapse:collapse;margin:1em 0}"
            "td,th{border:1px solid #ccc;padding:.3em .6em;text-align:left}th{background:#f3f4f7}"
            "blockquote{border-left:3px solid #ccc;margin-left:0;padding-left:1em;color:#555}"
            "pre{background:#f5f5f5;padding:.8em;overflow:auto}</style></head>\n"
            "<body>\n" + "\n".join(body) + "\n</body></html>\n")


def write_docx(blocks: Sequence[Block], target: Path) -> None:
    docx = require_docx()
    from docx.shared import Pt

    document = docx.Document()
    styles = {s.name for s in document.styles}

    def style(name: str, fallback: str = "Normal") -> str:
        return name if name in styles else fallback

    def fill(paragraph, runs: Sequence[Run]) -> None:
        for run in runs:
            piece = paragraph.add_run(run.text + (f" ({run.link})" if run.link and run.link != run.text else ""))
            piece.bold = run.bold or None
            piece.italic = run.italic or None
            if run.code:
                piece.font.name = "Consolas"

    for block in blocks:
        if block.kind == "heading":
            fill(document.add_heading(level=min(block.level, 9)), block.runs)
        elif block.kind in ("bullet", "number"):
            base = "List Bullet" if block.kind == "bullet" else "List Number"
            name = base if block.level == 0 else f"{base} {min(3, block.level + 1)}"
            fill(document.add_paragraph(style=style(name, style(base))), block.runs)
        elif block.kind == "quote":
            fill(document.add_paragraph(style=style("Quote")), block.runs)
        elif block.kind == "code":
            paragraph = document.add_paragraph(style=style("No Spacing"))
            piece = paragraph.add_run(block.text)
            piece.font.name = "Consolas"
            piece.font.size = Pt(9.5)
        elif block.kind == "table":
            width = max(len(row) for row in block.rows)
            table = document.add_table(rows=0, cols=width)
            if "Table Grid" in styles:
                table.style = "Table Grid"
            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for column, value in enumerate(row):
                    cells[column].text = value
                    if row_index == 0:
                        for run in cells[column].paragraphs[0].runs:
                            run.bold = True
        else:
            fill(document.add_paragraph(), block.runs)
    target.parent.mkdir(parents=True, exist_ok=True)
    scratch = target.with_name(target.name + ".part")
    try:
        document.save(str(scratch))
        scratch.replace(target)
    finally:
        if scratch.exists():
            scratch.unlink()


# ============================================================ Office -> PDF
NO_OFFICE_MESSAGE = (
    "Making a PDF needs Microsoft Office or LibreOffice on this computer: they draw the pages exactly "
    "as they look, Promak only asks them to. Install LibreOffice (free) and try again, or save as PDF "
    "from Word or Excel."
)


def find_libreoffice() -> Optional[str]:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    candidates = []
    if sys.platform.startswith("win"):
        for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                     os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            candidates.append(Path(base) / "LibreOffice" / "program" / "soffice.exe")
    elif sys.platform == "darwin":
        candidates.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))
    else:
        candidates += [Path("/usr/bin/soffice"), Path("/usr/lib/libreoffice/program/soffice"),
                       Path("/opt/libreoffice/program/soffice")]
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate)
    return None


def microsoft_office_apps() -> Tuple[str, ...]:
    """The Office programs Windows knows about ('Word', 'Excel', 'PowerPoint')."""
    if not sys.platform.startswith("win"):
        return ()
    try:
        import winreg
    except ImportError:  # pragma: no cover
        return ()
    found = []
    for app in ("Word", "Excel", "PowerPoint"):
        try:
            winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, f"{app}.Application"))
            found.append(app)
        except OSError:
            continue
    return tuple(found)


def office_pdf_converter() -> Optional[str]:
    """'LibreOffice', 'Microsoft Office' or None - what can make PDFs here."""
    if microsoft_office_apps():
        return "Microsoft Office"
    if find_libreoffice():
        return "LibreOffice"
    return None


_OFFICE_APP = {".doc": "Word", ".docx": "Word", ".odt": "Word", ".rtf": "Word",
               ".xls": "Excel", ".xlsx": "Excel", ".ods": "Excel", ".xlsm": "Excel",
               ".ppt": "PowerPoint", ".pptx": "PowerPoint", ".odp": "PowerPoint"}

_POWERSHELL = {
    "Word": ("$a = New-Object -ComObject Word.Application; $a.Visible = $false; "
             "$d = $a.Documents.Open($src, $false, $true); $d.SaveAs([ref] $dst, [ref] 17); "
             "$d.Close($false); $a.Quit()"),
    "Excel": ("$a = New-Object -ComObject Excel.Application; $a.Visible = $false; $a.DisplayAlerts = $false; "
              "$d = $a.Workbooks.Open($src, 0, $true); $d.ExportAsFixedFormat(0, $dst); "
              "$d.Close($false); $a.Quit()"),
    "PowerPoint": ("$a = New-Object -ComObject PowerPoint.Application; "
                   "$d = $a.Presentations.Open($src, $true, $false, $false); $d.SaveAs($dst, 32); "
                   "$d.Close(); $a.Quit()"),
}


def office_to_pdf(source: Path, target: Path, timeout: int = 300) -> str:
    """Ask Office or LibreOffice to print ``source`` to ``target``; returns who did it."""
    app = _OFFICE_APP.get(source.suffix.lower())
    if app and app in microsoft_office_apps():
        script = (f"$ErrorActionPreference = 'Stop'; $src = '{_ps(source.resolve())}'; "
                  f"$dst = '{_ps(target.resolve())}'; " + _POWERSHELL[app])
        result = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                                capture_output=True, text=True, errors="replace", timeout=timeout,
                                **SUBPROCESS_QUIET)
        if result.returncode == 0 and target.exists():
            return f"Microsoft {app}"
        log.warning("Office could not make the PDF: %s", result.stderr[-500:])
    soffice = find_libreoffice()
    if not soffice:
        raise ConvertError(NO_OFFICE_MESSAGE)
    with tempfile.TemporaryDirectory(prefix="promak-pdf-") as work:
        profile = Path(work) / "profile"
        command = [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore",
                   "--convert-to", "pdf", "--outdir", work, str(source)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, errors="replace",
                                    timeout=timeout, **SUBPROCESS_QUIET)
        except subprocess.TimeoutExpired as exc:
            raise ConvertError("LibreOffice took too long and was stopped.") from exc
        made = Path(work) / f"{source.stem}.pdf"
        if not made.exists():
            raise ConvertError("LibreOffice could not make the PDF: "
                               + ((result.stderr or result.stdout).strip().splitlines() or ["no reason given"])[-1])
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(made), str(target))
    return "LibreOffice"


def _ps(path: Path) -> str:
    return str(path).replace("'", "''")


# ================================================================ the engine
class ConvertBatch(BatchEngine):
    what = "file(s)"

    def __init__(self, options: ConvertOptions, **kwargs) -> None:
        super().__init__(**kwargs)
        self.options = options

    def process_one(self, job: FileJob, report) -> None:
        target = self.options.target
        suffix = job.source.suffix.lower()
        if suffix not in SOURCES[target]:
            allowed = ", ".join(e.lstrip(".").upper() for e in SOURCES[target])
            raise ConvertError(f"A {suffix.lstrip('.').upper()} file cannot become {target.upper()} here: "
                               f"that needs {allowed}.")
        report(10, "reading")
        if target == TO_PDF:
            out = output_path(job, self.options.suffix.strip(), ".pdf", self.options.overwrite)
            who = office_to_pdf(job.source, out)
            self._done(job, [out], f"PDF made by {who}")
        elif target == TO_CSV:
            self._to_csv(job)
        elif target == TO_XLSX:
            self._to_xlsx(job)
        else:
            self._document(job, target)

    def _done(self, job: FileJob, written: List[Path], message: str) -> None:
        job.output = written[0] if len(written) == 1 else written[0].parent
        job.output_bytes = sum(p.stat().st_size for p in written)
        job.info["result"] = human_size(job.output_bytes)
        job.message = message

    def _blocks(self, job: FileJob) -> List[Block]:
        suffix = job.source.suffix.lower()
        if suffix in WORD_EXTENSIONS:
            return read_docx(job.source)
        from promak.tools.text.engine import read_text

        text = read_text(job.source)
        return read_markdown(text) if suffix in (".md", ".markdown") else read_plain_text(text)

    def _document(self, job: FileJob, target: str) -> None:
        blocks = self._blocks(job)
        self.check_cancel()
        out = output_path(job, self.options.suffix.strip(), f".{target}", self.options.overwrite)
        if target == TO_DOCX:
            write_docx(blocks, out)
        else:
            if target == TO_TXT:
                body = to_text(blocks)
            elif target == TO_MD:
                body = to_markdown(blocks)
            else:
                body = to_html(blocks, job.source.stem)
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(body, encoding="utf-8")
        tables = sum(1 for b in blocks if b.kind == "table")
        self._done(job, [out], f"{len(blocks)} paragraph(s)" + (f", {tables} table(s)" if tables else ""))

    def _to_csv(self, job: FileJob) -> None:
        sheets = read_xlsx(job.source)
        if not self.options.every_sheet:
            sheets = dict(list(sheets.items())[:1])
        written = []
        for name, rows in sheets.items():
            self.check_cancel()
            ending = self.options.suffix.strip()
            if len(sheets) > 1:
                ending = f"{ending} - {safe_filename(name, fallback='sheet')}"
            out = output_path(job, ending, ".csv", self.options.overwrite)
            write_csv(out, rows, delimiter=self.options.delimiter)
            written.append(out)
        if not written:
            raise ConvertError("The workbook has no sheet.")
        self._done(job, written, f"{len(written)} CSV file(s)")

    def _to_xlsx(self, job: FileJob) -> None:
        rows, dialect = read_csv(job.source)
        comma = dialect.delimiter == ";"
        typed_rows = [[typed(cell, comma) if index else cell for cell in row] for index, row in enumerate(rows)]
        out = output_path(job, self.options.suffix.strip(), ".xlsx", self.options.overwrite)
        write_xlsx(out, {job.source.stem: typed_rows})
        self._done(job, [out], f"{len(rows)} row(s)")

    def describe_result(self, job: FileJob) -> str:
        return job.message or super().describe_result(job)

