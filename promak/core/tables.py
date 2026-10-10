"""Reading and writing tables: CSV files and Excel workbooks.

Shared by the document converter and the spreadsheet merger.  CSV files
come in many dialects - comma or semicolon (the usual one where the
decimal mark is a comma), UTF-8 or the old Windows encoding - so the
dialect is worked out from the file itself.  Excel files are read and
written with openpyxl (MIT).

Nothing here imports Qt.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from promak.core.imaging import ImageToolError

CSV_EXTENSIONS = (".csv", ".tsv")
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
TABLE_EXTENSIONS = CSV_EXTENSIONS + EXCEL_EXTENSIONS

Row = List[Any]


class TableError(ImageToolError):
    """A table that cannot be read or written, told in plain words."""


@dataclass
class CsvDialect:
    delimiter: str = ","
    encoding: str = "utf-8-sig"


def require_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - depends on the machine
        raise TableError("openpyxl is missing, so Excel files cannot be read or written. "
                         "Run install_windows.bat again, or:  pip install -U openpyxl") from exc
    return openpyxl


# --------------------------------------------------------------------- CSV
def _decode(data: bytes) -> tuple:
    for encoding in ("utf-8-sig", "utf-16") if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig",):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            pass
    return data.decode("cp1252", errors="replace"), "cp1252"


def read_csv(path: Path) -> tuple:
    """``(rows, CsvDialect)`` of a CSV file, its delimiter guessed."""
    try:
        data = Path(path).read_bytes()
    except OSError as exc:
        raise TableError(f"The file cannot be read: {exc}") from exc
    text, encoding = _decode(data)
    sample = text[:20000]
    if Path(path).suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
            delimiter = max(counts, key=counts.get) if any(counts.values()) else ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    return rows, CsvDialect(delimiter, encoding)


def write_csv(path: Path, rows: Sequence[Sequence[Any]], delimiter: str = ",", encoding: str = "utf-8-sig") -> None:
    """Write rows; dates as 2026-07-12, numbers with the decimal mark that suits the delimiter."""
    comma_decimal = delimiter == ";"
    path.parent.mkdir(parents=True, exist_ok=True)
    scratch = path.with_name(path.name + ".part")
    try:
        with open(scratch, "w", encoding=encoding, newline="") as handle:
            writer = csv.writer(handle, delimiter=delimiter)
            for row in rows:
                writer.writerow([cell_text(value, comma_decimal) for value in row])
        scratch.replace(path)
    finally:
        if scratch.exists():
            scratch.unlink()


def cell_text(value: Any, comma_decimal: bool = False) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if (value.hour or value.minute or value.second) \
            else value.strftime("%Y-%m-%d")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        text = repr(value) if value != int(value) or abs(value) >= 1e15 else str(int(value))
        return text.replace(".", ",") if comma_decimal else text
    return str(value)


_INT = re.compile(r"^-?\d{1,15}$")
_FLOAT_DOT = re.compile(r"^-?\d+\.\d+$")
_FLOAT_COMMA = re.compile(r"^-?\d+,\d+$")


def typed(value: str, comma_decimal: bool = False) -> Any:
    """A CSV cell as a number when it clearly is one; otherwise as it is.

    Numbers with leading zeros (postcodes, codes) stay text.
    """
    text = value.strip()
    if not text:
        return None
    if _INT.match(text) and not (len(text.lstrip("-")) > 1 and text.lstrip("-").startswith("0")):
        return int(text)
    if (_FLOAT_COMMA if comma_decimal else _FLOAT_DOT).match(text):
        return float(text.replace(",", "."))
    return value


# ------------------------------------------------------------------- Excel
def read_xlsx(path: Path) -> Dict[str, List[Row]]:
    """Every sheet of a workbook: ``{sheet name: rows}`` (values, not formulas)."""
    openpyxl = require_openpyxl()
    try:
        book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise TableError(f"This file cannot be read as an Excel workbook ({exc}).") from exc
    try:
        sheets: Dict[str, List[Row]] = {}
        for sheet in book.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            while rows and all(cell is None or str(cell).strip() == "" for cell in rows[-1]):
                rows.pop()
            width = max((i + 1 for row in rows for i, cell in enumerate(row) if cell not in (None, "")), default=0)
            sheets[sheet.title] = [row[:width] for row in rows]
        return sheets
    finally:
        book.close()


def write_xlsx(path: Path, sheets: Dict[str, Sequence[Sequence[Any]]], bold_header: bool = True) -> None:
    """Write one sheet per entry; the first row bold and frozen, columns sized."""
    openpyxl = require_openpyxl()
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(_sheet_name(name, book.sheetnames))
        widths: Dict[int, int] = {}
        for row in rows:
            sheet.append([_excel_value(value) for value in row])
            for index, value in enumerate(row, start=1):
                widths[index] = max(widths.get(index, 0), min(60, len(cell_text(value))))
        if bold_header and rows:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = max(8, width + 2)
    if not book.sheetnames:
        book.create_sheet("Sheet1")
    path.parent.mkdir(parents=True, exist_ok=True)
    scratch = path.with_name(path.name + ".part")
    try:
        book.save(str(scratch))
        scratch.replace(path)
    finally:
        if scratch.exists():
            scratch.unlink()


def _excel_value(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=",):
        return "'" + value  # never turn text into a formula
    return value


def _sheet_name(name: str, taken: Sequence[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", name or "Sheet").strip()[:31] or "Sheet"
    candidate, counter = clean, 2
    while candidate.lower() in (t.lower() for t in taken):
        suffix = f" ({counter})"
        candidate = clean[: 31 - len(suffix)] + suffix
        counter += 1
    return candidate


# ------------------------------------------------------------------- any
def read_table(path: Path) -> Dict[str, List[Row]]:
    """A CSV (one sheet named after the file) or every sheet of a workbook."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in CSV_EXTENSIONS:
        rows, dialect = read_csv(path)
        comma = dialect.delimiter == ";"
        return {path.stem: [[typed(cell, comma) for cell in row] for row in rows]}
    if suffix in EXCEL_EXTENSIONS:
        return read_xlsx(path)
    raise TableError(f"'{path.name}' is not a CSV or Excel file.")


def first_sheet(sheets: Dict[str, List[Row]], name: Optional[str] = None) -> List[Row]:
    if name and name in sheets:
        return sheets[name]
    return next(iter(sheets.values()), [])
