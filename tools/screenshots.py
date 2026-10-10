#!/usr/bin/env python3
"""Take pictures of every Promak screen, without showing a window.

    python tools/screenshots.py out/            every tool, light and dark
    python tools/screenshots.py out/ --size 1600x900

Qt runs in its "offscreen" mode, so this works on a server with no screen.
Each queue is filled with a few made-up rows and log lines first, so the
pictures show the screens the way they look while in use.  Nothing is
downloaded and no file is converted.  The settings are written to a
temporary folder, never to the real ones.

Used by the "screens" job of the CI: the pictures end up on the
``ci-renders`` branch, where a session in the cloud can look at them.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_SANDBOX = Path(tempfile.mkdtemp(prefix="promak-shots-"))
os.environ["APPDATA"] = str(_SANDBOX / "appdata")
os.environ["XDG_DATA_HOME"] = str(_SANDBOX / "appdata")


def _sample_pictures(folder: Path) -> list:
    """Three small pictures to put in the queue of the picture tools."""
    folder.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        return []
    paths = []
    for index, (name, colour) in enumerate(
        (("holiday-photo.jpg", (52, 120, 200)), ("company-logo.png", (230, 120, 40)),
         ("screenshot.png", (40, 170, 120)))
    ):
        image = Image.new("RGB", (640, 420), (245, 245, 245))
        draw = ImageDraw.Draw(image)
        draw.ellipse((120 + index * 30, 60, 420 + index * 30, 360), fill=colour)
        draw.rectangle((40, 300, 600, 380), fill=(30, 30, 30))
        path = folder / name
        image.save(path)
        paths.append(path)
    return paths


def _sample_documents(folder: Path) -> list:
    """A few small documents for the document tools."""
    folder.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        return []
    paths = []
    for name, pages in (("Annual report.pdf", 12), ("Scanned contract.pdf", 3)):
        sheets = []
        for number in range(pages):
            sheet = Image.new("RGB", (620, 877), "white")
            draw = ImageDraw.Draw(sheet)
            draw.rectangle((60, 60, 560, 110), fill=(47, 91, 234))
            for line in range(18):
                draw.rectangle((60, 150 + line * 34, 520 - (line % 4) * 40, 162 + line * 34), fill=(200, 205, 215))
            draw.text((60, 830), f"{number + 1}", fill="black")
            sheets.append(sheet)
        path = folder / name
        sheets[0].save(path, "PDF", save_all=True, append_images=sheets[1:])
        paths.append(path)
    notes = folder / "Meeting notes.md"
    notes.write_text("# Meeting\n\n- Budget\n- Dates\n", encoding="utf-8")
    prices = folder / "Price list.csv"
    prices.write_text("Item;Price\nChair;49,90\nTable;129,00\n", encoding="utf-8")
    paths += [notes, prices]
    try:
        from promak.core.tables import write_xlsx

        write_xlsx(folder / "Budget 2026.xlsx", {"Budget": [["Month", "Spent"], ["January", 1200]]})
        paths.append(folder / "Budget 2026.xlsx")
        import docx

        document = docx.Document()
        document.add_heading("Offer", 1)
        document.add_paragraph("Thank you for your request.")
        document.save(str(folder / "Offer.docx"))
        paths.append(folder / "Offer.docx")
    except Exception:  # an optional component is missing: fewer samples
        pass
    return paths


def _fill_video(panel, destination: Path) -> None:
    from promak.tools.video.models import Stage

    panel.destination_input.setText(str(destination))
    panel._append_jobs(
        [
            "https://a-video-site.example/watch?v=first",
            "https://a-news-site.example/article-with-video",
            "https://a-video-site.example/watch?v=third",
        ],
        destination,
        titles=["How bridges are built", "Evening news, 7 October", "A short talk on gears"],
    )
    states = [(Stage.DONE, 100), (Stage.TRANSCRIBE, 64), (Stage.QUEUED, 0)]
    for job, (stage, overall) in zip(panel._jobs, states):
        detail = {100: "Saved", 64: "Transcribing 40%", 0: ""}[overall]
        panel._apply_job_state(job.id, stage.value, overall, detail, job.display_name)
    for line in (
        "Added 3 link(s) to the queue.",
        "How bridges are built: downloaded (182 MB).",
        "How bridges are built: MP3 written.",
        "How bridges are built: transcript written.",
        "Evening news, 7 October: transcribing...",
    ):
        panel._log("info", line)


def _fill_files(panel, pictures, destination: Path) -> None:
    panel.destination_input.setText(str(destination))
    if pictures:
        panel.add_files(pictures)
    panel._log("info", "Ready: select a row to see its preview.")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("out", type=Path)
    parser.add_argument("--size", default="1920x1010", help="window size, WIDTHxHEIGHT")
    args = parser.parse_args(argv)
    width, height = (int(v) for v in args.size.lower().split("x"))
    args.out.mkdir(parents=True, exist_ok=True)

    from PySide6.QtWidgets import QApplication

    from promak.ui.main_window import MainWindow
    from promak.ui.theme import stylesheet

    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(stylesheet("light"))
    window = MainWindow()
    window.resize(width, height)
    window.show()
    app.processEvents()

    destination = _SANDBOX / "Downloads" / "Promak"
    pictures = _sample_pictures(_SANDBOX / "pictures") + _sample_documents(_SANDBOX / "documents")
    for index in range(window.stack.count()):
        page = window.stack.widget(index)
        if hasattr(page, "screenshot_sample"):
            # screens that are not a plain queue fill themselves
            page.screenshot_sample(_SANDBOX)
        elif hasattr(page, "url_input"):
            _fill_video(page, destination)
        elif hasattr(page, "set_folder"):
            albums = _SANDBOX / "albums"
            for name in ("Summer by the sea", "Grandma's birthday", "07 - School play", "Winter 2"):
                (albums / name).mkdir(parents=True, exist_ok=True)
            page.set_folder(albums)
        elif hasattr(page, "add_files"):
            accepted = tuple(getattr(page, "ACCEPTED_EXTENSIONS", ()))
            files = [p for p in pictures if p.suffix.lower() in accepted]
            # documents first, so a document tool does not look like a picture tool
            files.sort(key=lambda p: p.suffix.lower() in (".png", ".jpg"))
            _fill_files(page, files, destination)

    taken = []
    for theme in ("light", "dark"):
        window.apply_theme(theme)
        for row in window.tool_rows():
            window.tool_list.setCurrentRow(row)
            tool_id = window.tool_list.item(row).data(256)  # Qt.UserRole
            for _ in range(3):
                app.processEvents()
            path = args.out / f"{tool_id}-{theme}.png"
            window.grab().save(str(path))
            taken.append(path)
    # the top of the sidebar, enlarged, to check the logo and the switch
    window.apply_theme("light")
    app.processEvents()
    corner = window.grab().copy(0, 0, 300, 140)
    corner = corner.scaled(corner.width() * 2, corner.height() * 2)
    corner.save(str(args.out / "sidebar-top-light.png"))
    taken.append(args.out / "sidebar-top-light.png")

    for path in taken:
        print(path)
    window.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
